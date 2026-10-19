"""leitura das células de entrada."""

from __future__ import annotations

import hashlib
import math
import re
from pathlib import Path
from typing import Any

import openpyxl


class DataError(ValueError):
    pass


ABA_RECEITAS = 'Detalhamento das Receitas'
ABA_DESPESAS = 'Despesas por Secretaria'
ABA_FOLHA = 'Folha e Limites LRF'
ABA_DIVIDA = 'Dívida Ativa e Inadimplência'
ABA_BALANCO = 'Balanço e Limites LRF'

REC_ID, REC_CATEGORIA, REC_ORIGEM, REC_PREVISTO, REC_ARRECADADO = 0, 1, 2, 3, 4
DESP_ID, DESP_NOME, DESP_ORCADO, DESP_LIQUIDADA, DESP_FOLHA, DESP_INVEST = 0, 1, 3, 4, 5, 6
DIV_NOME, DIV_ESTOQUE, DIV_ESTAGIO, DIV_RECUPERACAO, DIV_PRESCRICAO, DIV_RENUNCIA = 0, 1, 2, 3, 4, 5
FOLHA_MATRICULA, FOLHA_CATEGORIA, FOLHA_VINCULO, FOLHA_LOTACAO, FOLHA_VENCIMENTO, FOLHA_FALTAS = 0, 1, 2, 3, 4, 7

LINHAS_RECEITAS = (5, 18)
LINHAS_DESPESAS = (5, 13)
LINHAS_DIVIDA = (5, 9)
PRIMEIRA_LINHA_FOLHA = 5

ENCARGOS_PATRONAIS = 1.22   # encargos da planilha
DIAS_FALTA_180 = 180 * 0.714

VINCULO_COMISSIONADO = 'Comissionado (DAS)'
CATEGORIA_ADMINISTRATIVA = 'Assistente Administrativo / Operacional'


def money(value: Any, label: str) -> float:
    """exige número finito e não negativo; rejeita fórmulas."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise DataError(f'{label}: esperado número de entrada; encontrado {value!r}.')
    if not math.isfinite(value) or value < 0:
        raise DataError(f'{label}: valor deve ser finito e não negativo.')
    return round(float(value), 2)


def ref(sheet: str, cell: str) -> str:
    return f"'{sheet}'!{cell}"


def ler_receitas(aba) -> list[dict]:
    receitas = []
    for linha in aba.iter_rows(min_row=LINHAS_RECEITAS[0], max_row=LINHAS_RECEITAS[1]):
        n = linha[0].row
        if not linha[REC_ID].value or not linha[REC_CATEGORIA].value:
            raise DataError(f'Receitas: identificador ou categoria ausente na linha {n}.')
        previsto = money(linha[REC_PREVISTO].value, f'Receitas!D{n}')
        arrecadado = money(linha[REC_ARRECADADO].value, f'Receitas!E{n}')
        receitas.append(dict(
            id=linha[REC_ID].value, name=linha[REC_CATEGORIA].value, origin=linha[REC_ORIGEM].value,
            planned=previsto, actual=arrecadado, gap=round(previsto - arrecadado, 2),
            execution=arrecadado / previsto if previsto else None,
            source=ref(aba.title, f'D{n}:E{n}')))
    return receitas


def ler_despesas(aba) -> list[dict]:
    secretarias = []
    for linha in aba.iter_rows(min_row=LINHAS_DESPESAS[0], max_row=LINHAS_DESPESAS[1]):
        n = linha[0].row
        liquidada = money(linha[DESP_LIQUIDADA].value, f'Despesas!{linha[DESP_LIQUIDADA].coordinate}')
        folha = money(linha[DESP_FOLHA].value, f'Despesas!{linha[DESP_FOLHA].coordinate}')
        investimento = money(linha[DESP_INVEST].value, f'Despesas!{linha[DESP_INVEST].coordinate}')
        secretarias.append(dict(
            id=linha[DESP_ID].value, name=linha[DESP_NOME].value,
            budget=money(linha[DESP_ORCADO].value, f'Despesas!D{n}'),
            actual=liquidada, payroll=folha, investment=investimento,
            # custeio aproximado, sem folha e investimentos
            eligible=round(max(0, liquidada - folha - investimento), 2),
            source=ref(aba.title, f'D{n}:G{n}')))
    return secretarias


def ler_divida(aba) -> list[dict]:
    dividas = []
    for linha in aba.iter_rows(min_row=LINHAS_DIVIDA[0], max_row=LINHAS_DIVIDA[1]):
        n = linha[0].row
        colunas = (DIV_ESTOQUE, DIV_RECUPERACAO, DIV_PRESCRICAO, DIV_RENUNCIA)
        estoque, recuperacao, perda, renuncia = (
            money(linha[c].value, f'Divida!{linha[c].coordinate}') for c in colunas)
        if recuperacao > estoque:
            raise DataError(f'Divida!D{n}: recuperação estimada maior que o estoque.')
        dividas.append(dict(
            id=f'DIV-{n}', name=linha[DIV_NOME].value, stock=estoque, recovery=recuperacao,
            reported_loss=perda, waivers=renuncia, stage=linha[DIV_ESTAGIO].value,
            source=ref(aba.title, f'A{n}:F{n}')))
    return dividas


def ler_folha(aba) -> dict:
    """agrega servidores e vencimentos."""
    quantidade = 0
    soma_vencimentos = 0
    dias_de_falta = 0
    por_unidade = {}
    por_grupo = {}
    for linha in aba.iter_rows(min_row=PRIMEIRA_LINHA_FOLHA, values_only=True):
        matricula = linha[FOLHA_MATRICULA]
        if not isinstance(matricula, str) or not matricula.startswith('SERV-'):
            continue
        quantidade += 1
        soma_vencimentos += money(linha[FOLHA_VENCIMENTO], f'Folha: vencimento de {matricula}')
        dias_de_falta += money(linha[FOLHA_FALTAS], f'Folha: faltas de {matricula}')
        unidade = str(linha[FOLHA_LOTACAO])
        por_unidade[unidade] = por_unidade.get(unidade, 0) + 1
        grupo = por_grupo.setdefault((str(linha[FOLHA_CATEGORIA]), str(linha[FOLHA_VINCULO])), dict(count=0, salary_sum=0))
        grupo['count'] += 1
        grupo['salary_sum'] += money(linha[FOLHA_VENCIMENTO], f'Folha: vencimento de {matricula}')
    if not quantidade:
        raise DataError('Não foi encontrado nenhum servidor na aba de folha.')
    grupos = [dict(category=c, link=v, count=g['count'], monthly_cost=round(g['salary_sum'] * ENCARGOS_PATRONAIS, 2))
              for (c, v), g in por_grupo.items()]
    grupos.sort(key=lambda g: g['monthly_cost'], reverse=True)
    return dict(count=quantidade, salary_sum=soma_vencimentos, absence_days=dias_de_falta,
                by_unit=por_unidade, groups=grupos)


def montar_avisos() -> list[dict]:
    """limitações da planilha."""
    return [
        dict(id='periodo', title='Período e base fiscal precisam ser confirmados',
             text='As receitas dizem 12 meses, mas as despesas não informam expressamente o período. A simulação assume a mesma janela de 12 meses. Fica como limitação do estudo; o período exato deve ser confirmado na implementação.',
             source=ref(ABA_DESPESAS, 'E4:E13')),
        dict(id='rcl', title='Receita total não é RCL apurada',
             text='A aba Balanço referencia a soma das receitas como RCL. O sistema mostra apenas a razão simples folha/receita, sem emitir conformidade com a LRF.',
             source=ref(ABA_BALANCO, 'C5:C6')),
        dict(id='primario', title='Saldo simples, sem apuração de resultado primário',
             text='Receitas menos despesas inclui rubricas que exigem classificação e ajustes. O saldo exibido não comprova resultado primário nem disponibilidade de caixa.',
             source=ref(ABA_BALANCO, 'C10')),
        dict(id='rotulos', title='Rótulos de conformidade não foram validados',
             text='Os avisos da aba Balanço são textos fixos. As bases para saúde, educação, FUNDEB e restos a pagar precisam de documentação e revisão do responsável.',
             source=ref(ABA_BALANCO, 'C7:E11')),
        dict(id='divida', title='Recuperação de dívida é uma hipótese',
             text='A estimativa do arquivo não tem prazo de realização. Ela não é dinheiro em caixa. Perdas informadas, renúncias e diferenças de arrecadação não são somadas ao estoque.',
             source=ref(ABA_DIVIDA, 'B5:F9')),
        dict(id='base_cortes', title='Base de economia é aproximada',
             text='Despesa liquidada menos folha e investimentos pressupõe componentes comparáveis, contidos no total e sem sobreposição. Não há detalhamento de contratos nem de serviços essenciais.',
             source=ref(ABA_DESPESAS, 'E5:G13')),
    ]


def nome_curto(nome: str) -> str:
    """remove o detalhamento da rubrica."""
    return re.sub(r'\s*\([^)]*\)', '', nome.split(' - ')[0])


def percentual(valor: float) -> str:
    return f'{valor * 100:.0f}%'.replace('.', ',')


def montar_situacao(metrics: dict, receitas: list, secretarias: list, dividas: list) -> list[dict]:
    """diagnóstico com dados da planilha."""
    def item(label, value, kind='money'):
        return dict(label=label, value=value, kind=kind)

    achados = []

    achados.append(dict(
        id='resultado', title='Resultado do período', value=metrics['balance'], kind='money',
        severity='alerta' if metrics['balance'] < 0 else 'info',
        note='Receitas arrecadadas menos despesas liquidadas. É um saldo simples: não é resultado primário nem caixa.',
        items=[item('Receita arrecadada', metrics['revenue']), item('Despesa liquidada', metrics['expenses'])],
        source=ref(ABA_BALANCO, 'C10')))

    faltas = sorted((r for r in receitas if r['gap'] > 0), key=lambda r: r['gap'], reverse=True)[:4]
    falta_total = round(metrics['planned_revenue'] - metrics['revenue'], 2)
    achados.append(dict(
        id='receita_abaixo_do_previsto', title='Arrecadação abaixo do previsto', value=falta_total, kind='money',
        severity='alerta' if falta_total > 0 else 'info',
        note=f'A arrecadação ficou em {percentual(metrics["revenue"] / metrics["planned_revenue"])} do previsto. Maiores diferenças:',
        items=[item(f'{nome_curto(r["name"])}: {percentual(r["execution"])} do previsto', r['gap']) for r in faltas],
        source=ref(ABA_RECEITAS, 'D5:E18')))

    acima = [s for s in secretarias if s['actual'] > s['budget']]
    abaixo = sorted((s for s in secretarias if s['actual'] < s['budget']), key=lambda s: s['budget'] - s['actual'], reverse=True)[:2]
    achados.append(dict(
        id='despesa_acima_da_dotacao', title='Despesa acima da dotação', value=round(sum(s['actual'] - s['budget'] for s in acima), 2),
        kind='money', severity='alerta' if acima else 'info',
        note='Secretarias que gastaram mais que a dotação inicial; abaixo, as que ficaram mais abaixo do orçado.',
        items=[item(s['name'] + ' (acima)', round(s['actual'] - s['budget'], 2)) for s in acima]
              + [item(s['name'] + ' (abaixo)', round(s['budget'] - s['actual'], 2)) for s in abaixo],
        source=ref(ABA_DESPESAS, 'D5:E13')))

    peso = metrics['payroll_ratio_proxy']
    achados.append(dict(
        id='peso_da_folha', title='Peso da folha', value=peso, kind='percent',
        severity='atencao',
        note='Folha das secretarias dividida pela receita arrecadada. É uma razão simples, não o teste legal da LRF, que depende da RCL oficial.',
        items=[item('Folha das secretarias (ano)', metrics['payroll_departments']),
               item(f'Cargos de confiança ({metrics["das_count"]}), custo anual', round(metrics['das_monthly_cost'] * 12, 2)),
               item('Faltas e licenças (estimativa da planilha)', metrics['absence_proxy'], 'percent')],
        source=ref(ABA_DESPESAS, 'F5:F13')))

    renuncia = round(sum(d['waivers'] for d in dividas), 2)
    achados.append(dict(
        id='divida_ativa', title='Dívida ativa', value=metrics['debt_stock'], kind='money', severity='alerta',
        note=f'Só {percentual(metrics["estimated_recovery"] / metrics["debt_stock"])} do estoque tem recuperação estimada, sem prazo. Estoque não é dinheiro em caixa.',
        items=[item('Recuperação estimada', metrics['estimated_recovery']),
               item('Perda por prescrição já informada', metrics['reported_prescription_loss']),
               item('Renúncia fiscal e isenções', renuncia)],
        source=ref(ABA_DIVIDA, 'B5:F9')))

    convenios = next((r for r in receitas if r['name'].startswith('Conv')), None)
    infra = next((s for s in secretarias if s['id'] == 'SEC-INF'), None)
    if convenios and infra:
        achados.append(dict(
            id='convenios_e_obras', title='Convênios e obras', value=convenios['gap'], kind='money',
            severity='alerta' if convenios['execution'] < 0.5 else 'atencao',
            note=f'Os convênios arrecadaram {percentual(convenios["execution"])} do previsto, e Infraestrutura gastou menos que o orçado. Combina com o relato do cliente sobre repasses sem controle.',
            items=[item('Convênios: diferença para o previsto', convenios['gap']),
                   item('Infraestrutura: abaixo da dotação', round(infra['budget'] - infra['actual'], 2))],
            source=ref(ABA_RECEITAS, 'D16:E16')))

    achados.append(dict(
        id='restos_a_pagar', title='Restos a pagar sem cobertura', value=metrics['uncovered_payables'], kind='money',
        severity='atencao',
        note='Valor informado na planilha. A conformidade com o art. 42 da LRF depende de apuração oficial; o app não conclui sobre isso.',
        items=[], source=ref(ABA_BALANCO, 'C11')))
    return achados


def load_dataset(path: Path) -> dict:
    if not path.is_file():
        raise DataError('Planilha não encontrada na pasta data/. Veja o README.')
    planilha = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        abas = [ABA_RECEITAS, ABA_DESPESAS, ABA_FOLHA, ABA_DIVIDA, ABA_BALANCO]
        if not set(abas).issubset(planilha.sheetnames):
            raise DataError('A planilha não tem as cinco abas esperadas para este estudo de caso.')

        receitas = ler_receitas(planilha[ABA_RECEITAS])
        secretarias = ler_despesas(planilha[ABA_DESPESAS])
        dividas = ler_divida(planilha[ABA_DIVIDA])
        folha = ler_folha(planilha[ABA_FOLHA])
        balanco = planilha[ABA_BALANCO]

        if len({r['id'] for r in receitas}) != len(receitas) or len({s['id'] for s in secretarias}) != len(secretarias):
            raise DataError('Foram encontrados identificadores duplicados nas tabelas de receitas/despesas.')

        total_receitas = round(sum(r['actual'] for r in receitas), 2)
        total_despesas = round(sum(s['actual'] for s in secretarias), 2)
        folha_secretarias = round(sum(s['payroll'] for s in secretarias), 2)
        folha_mensal = round(folha['salary_sum'] * ENCARGOS_PATRONAIS, 2)
        comissionados = next((g for g in folha['groups'] if g['link'] == VINCULO_COMISSIONADO), dict(count=0, monthly_cost=0))
        administrativos = next((g for g in folha['groups'] if g['category'] == CATEGORIA_ADMINISTRATIVA), dict(count=0, monthly_cost=0))

        metrics = dict(
            revenue=total_receitas, expenses=total_despesas, balance=round(total_receitas - total_despesas, 2),
            planned_revenue=round(sum(r['planned'] for r in receitas), 2),
            payroll_departments=folha_secretarias,
            payroll_ratio_proxy=folha_secretarias / total_receitas if total_receitas else None,
            monthly_payroll=folha_mensal, payroll_times_12=round(folha_mensal * 12, 2),
            staff_count=folha['count'], absence_days=folha['absence_days'],
            das_count=comissionados['count'], das_monthly_cost=comissionados['monthly_cost'],
            admin_staff_count=administrativos['count'], admin_staff_monthly_cost=administrativos['monthly_cost'],
            absence_proxy=folha['absence_days'] / (folha['count'] * DIAS_FALTA_180),
            debt_stock=round(sum(d['stock'] for d in dividas), 2),
            estimated_recovery=round(sum(d['recovery'] for d in dividas), 2),
            reported_prescription_loss=round(sum(d['reported_loss'] for d in dividas), 2),
            uncovered_payables=money(balanco['C11'].value, 'Balanco!C11'))

        avisos = montar_avisos()
        avisos.append(dict(
            id='folha', title='Folha detalhada e agregado precisam ser reconciliados',
            text='Mensal vezes 12 não inclui 13º, férias ou outros ajustes. Não substitui o agregado de pessoal das secretarias. Absenteísmo reproduz a hipótese de 180 × 0,714 dias da planilha.',
            source=ref(ABA_FOLHA, f'E5:I{4 + folha["count"]}')))
        if any(s['payroll'] + s['investment'] > s['actual'] for s in secretarias):
            avisos.append(dict(
                id='sobreposicao', title='Componentes excedem a despesa',
                text='Alguma secretaria possui folha mais investimento maior que a despesa total. Sua base de economia foi limitada a zero e exige esclarecimento.',
                source=ref(ABA_DESPESAS, 'E5:G13')))

        return dict(
            filename=path.name, fingerprint=hashlib.sha256(path.read_bytes()).hexdigest(),
            metrics=metrics, situation=montar_situacao(metrics, receitas, secretarias, dividas),
            revenues=receitas, departments=secretarias, debts=dividas,
            staff_by_unit=folha['by_unit'], payroll_groups=folha['groups'], warnings=avisos,
            metric_sources=dict(
                revenue=ref(ABA_RECEITAS, 'E5:E18'), expenses=ref(ABA_DESPESAS, 'E5:E13'),
                payroll_departments=ref(ABA_DESPESAS, 'F5:F13'), debt_stock=ref(ABA_DIVIDA, 'B5:B9'),
                estimated_recovery=ref(ABA_DIVIDA, 'D5:D9'), payroll_groups=ref(ABA_FOLHA, f'B5:E{4 + folha["count"]}')))
    finally:
        planilha.close()
